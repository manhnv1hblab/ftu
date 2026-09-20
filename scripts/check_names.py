import openpyxl
import sys
from pathlib import Path
from difflib import get_close_matches

sys.stdout.reconfigure(encoding='utf-8')
DOCUMENT_DIR = Path(__file__).resolve().parents[1] / 'document'

# Partners
wb_part = openpyxl.load_workbook(DOCUMENT_DIR / '[CTTĐ_S27] Danh sách trường đối tác trao đổi.xlsx', data_only=True)
s_part = wb_part['Danh sách đối tác S27']
partners = []
for r in s_part.iter_rows(min_row=2, values_only=True):
    if r[3]:
        partners.append(str(r[3]).strip())

# Equivalence
wb_eq = openpyxl.load_workbook(DOCUMENT_DIR / 'Danh sách học phần tương đương (với các trường đối tác).xlsx', data_only=True)
s_eq = wb_eq['Tổng']
eq_unis = set()
for r in s_eq.iter_rows(min_row=2, values_only=True):
    if r[2]:
        eq_unis.add(str(r[2]).strip())

print(f'Total partners in S27: {len(partners)}')
print(f'Total unique unis in Tong: {len(eq_unis)}')

exact_matches = [p for p in partners if p in eq_unis]
print(f'Exact matches: {len(exact_matches)}')

no_matches = [p for p in partners if p not in eq_unis]
print(f'No exact match count: {len(no_matches)}')
for p in no_matches:
    close = get_close_matches(p, eq_unis, n=2, cutoff=0.5)
    print(f'  Partner: "{p}" -> Candidate in Tong: {close}')
